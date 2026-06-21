"""
SCRIPT 1 of 3 — extract_country_data.py

Extracts raw 5-year age-bracket population data for every country/area
from the official UN WPP2024 file POP/F02-1, into a flat CSV. Country-level
only — no regional grouping happens here (that's script 2's job).

Source file (confirmed URL):
https://population.un.org/wpp/assets/Excel%20Files/1_Indicator%20(Standard)/EXCEL_FILES/2_Population/WPP2024_POP_F02_1_POPULATION_5-YEAR_AGE_GROUPS_BOTH_SEXES.xlsx

Usage:
    python3 extract_country_data.py [path_to_xlsx] [year]
    python3 extract_country_data.py --download [year]   # fetch it first

Output:
    country_data_<year>.csv — one row per country/area with every 5-year
    age bracket (0-4 ... 100+), plus convenience columns for each of the
    cutoffs we report on: TotalPop, Under40, Under20, Age0to4, Over65
    (all in thousands, the file's native unit, plus a "_pct" share of
    that country's own population for each cutoff).

On the Under20 cutoff:
    The source file only has 5-year brackets, so there's no column that
    cuts cleanly at 18 the way there is at 40 (35-39 | 40-44) or 65
    (60-64 | 65-69) — 18 falls in the middle of the "15-19" bracket
    (ages 15-19 inclusive). Rather than interpolate, Under20_thousands
    sums the bracket through "15-19" as-is: (0-4)+(5-9)+(10-14)+(15-19),
    which is an EXACT figure for "ages 0-19" / "under 20" — no estimation
    involved.
"""

import sys
import urllib.request

import openpyxl

SOURCE_URL = (
    "https://population.un.org/wpp/assets/Excel%20Files/"
    "1_Indicator%20(Standard)/EXCEL_FILES/2_Population/"
    "WPP2024_POP_F02_1_POPULATION_5-YEAR_AGE_GROUPS_BOTH_SEXES.xlsx"
)
DEFAULT_FILENAME = (
    "WPP2024_POP_F02_1_POPULATION_5-YEAR_AGE_GROUPS_BOTH_SEXES.xlsx"
)
DEFAULT_YEAR = 2025

# Header columns (1-indexed) shared by the "Estimates" and variant sheets:
#   1 Index | 2 Variant | 3 Location | 4 Notes | 5 Location code
#   6 ISO3 | 7 ISO2 | 8 SDMX code | 9 Type | 10 Parent code | 11 Year
#   12..32 = age brackets: 0-4, 5-9, ..., 95-99, 100+  (21 brackets)
AGE_LABELS = [
    "0-4", "5-9", "10-14", "15-19", "20-24", "25-29", "30-34", "35-39",
    "40-44", "45-49", "50-54", "55-59", "60-64", "65-69", "70-74", "75-79",
    "80-84", "85-89", "90-94", "95-99", "100+",
]
AGE_COL_START = 12        # 1-indexed column of "0-4"
UNDER_40_BRACKETS = 8     # 0-4 through 35-39 (clean boundary: 40 starts 40-44)
UNDER_20_BRACKETS = 4     # 0-4 through 15-19 (clean boundary: 20 starts 20-24)
OVER_65_START_BRACKET = 13   # index of "65-69" (clean boundary)


def maybe_download(target_path: str) -> None:
    print(f"Downloading {SOURCE_URL} ...", file=sys.stderr)
    print("(this file is ~60MB, may take a few minutes)", file=sys.stderr)
    urllib.request.urlretrieve(SOURCE_URL, target_path)


def pick_sheet(year: int) -> str:
    # "Estimates" covers 1950-2023 (observed). Everything 2024+ lives in
    # the variant sheets; Medium is the standard UN central scenario.
    return "Estimates" if year <= 2023 else "Medium variant"


def main():
    args = sys.argv[1:]
    download = "--download" in args
    args = [a for a in args if a != "--download"]

    path = DEFAULT_FILENAME
    year = DEFAULT_YEAR
    if download:
        if args:
            year = int(args[0])
        maybe_download(path)
    else:
        if args:
            path = args[0]
        if len(args) > 1:
            year = int(args[1])

    sheet_name = pick_sheet(year)
    print(f"Reading sheet '{sheet_name}', year {year}, from {path} ...",
          file=sys.stderr)

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name]

    out_rows = []
    # ws.max_row is unreliable on this file (bad cached dimensions in the
    # UN's own export), so force an explicit generous range instead.
    for row in ws.iter_rows(min_row=18, max_row=500000,
                             min_col=1, max_col=AGE_COL_START + 21,
                             values_only=True):
        if row[0] is None or row[10] != year or row[8] != "Country/Area":
            continue

        ages = row[AGE_COL_START - 1: AGE_COL_START - 1 + 21]
        if any(v is None for v in ages):
            continue  # incomplete row (shouldn't happen for Country/Area)

        total = sum(ages)
        under40 = sum(ages[:UNDER_40_BRACKETS])
        age0to4 = ages[0]
        over65 = sum(ages[OVER_65_START_BRACKET:])
        under20 = sum(ages[:UNDER_20_BRACKETS])

        def pct(part):
            return round(100 * part / total, 2) if total else None

        out_rows.append({
            "Country": row[2],
            "ISO3": row[5],
            "LocationCode": row[4],
            "Year": year,
            **{label: round(val, 2) for label, val in zip(AGE_LABELS, ages)},
            "TotalPop_thousands": round(total, 2),
            "Under40_thousands": round(under40, 2),
            "Under40_pct": pct(under40),
            "Under20_thousands": round(under20, 2),
            "Under20_pct": pct(under20),
            "Age0to4_thousands": round(age0to4, 2),
            "Age0to4_pct": pct(age0to4),
            "Over65_thousands": round(over65, 2),
            "Over65_pct": pct(over65),
        })

    out_rows.sort(key=lambda d: d["TotalPop_thousands"], reverse=True)

    out_path = f"country_data_{year}.csv"
    fieldnames = (
        ["Country", "ISO3", "LocationCode", "Year"] + AGE_LABELS
        + ["TotalPop_thousands",
           "Under40_thousands", "Under40_pct",
           "Under20_thousands", "Under20_pct",
           "Age0to4_thousands", "Age0to4_pct",
           "Over65_thousands", "Over65_pct"]
    )

    import csv
    with open(out_path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        w.writerows(out_rows)

    print(f"\nWrote {len(out_rows)} countries/areas to {out_path}")


if __name__ == "__main__":
    main()
